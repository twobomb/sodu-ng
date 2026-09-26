# -*- coding: utf-8 -*-
import zipfile, re
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

z = zipfile.ZipFile('_out.xlsx')
W = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
sheet = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))

# 1) Никакого rich text в листе
raw = z.read('xl/worksheets/sheet1.xml').decode('utf-8')
print('rich <r> runs on sheet:', raw.count('<r>'))
print('sharedStrings.xml present:', 'xl/sharedStrings.xml' in z.namelist())

# 2) Невалидные inline-строки без <is>
bad = 0
for c_ in sheet.iter(W + 'c'):
    if c_.get('t') == 'inlineStr' and c_.find(W + 'is') is None:
        bad += 1
print('empty inlineStr without <is>:', bad)

# 3) t="s" без <v>
sbad = 0
for c_ in sheet.iter(W + 'c'):
    if c_.get('t') == 's' and c_.find(W + 'v') is None:
        sbad += 1
print('t="s" without v:', sbad)

# 4) E-техника (подразделение+тип) и D-переносы
wb = load_workbook('_out.xlsx')
ws = wb.active
print('--- E sample ---')
ee = [ws.cell(r, 5).value for r in range(5, ws.max_row + 1) if ws.cell(r, 5).value]
print(ee[:6])
print('--- D sample (headers/blank lines) ---')
for r in range(5, ws.max_row + 1):
    v = ws.cell(r, 4).value
    if isinstance(v, str) and v:
        print('row', r, repr(v[:60]), '| blank line present:', '\n\n' in v)
        break