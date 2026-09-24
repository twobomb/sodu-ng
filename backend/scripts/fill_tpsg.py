# -*- coding: utf-8 -*-
"""Заполняет шаблон ТПСГ (.xlsx) данными через openpyxl.
Полностью пересобирает область данных и нижнюю шапку, сохраняя стили шаблона.

Usage:
  python fill_tpsg.py <template.xlsx> <data.json> <out.xlsx>

data.json:
{
  "date": "ДД.ММ.ГГГГ",
  "officers": { "nach": "", "st": "", "pom": "", "disp": "" },
  "data_rows": [
    { "type": "garrison", "text": "<название гарнизона>", "values": {"3": 2, ...} },
    { "type": "dept", "text": "<имя>", "col2": "<вид>", "values": {...} },
    { "type": "subtotal", "text": "Итого за <г>", "values": {...} },
    { "type": "grand", "text": "Итого за ТПСГ...", "values": {...} }
  ]
}
"""
import sys
import json
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

RED_FILL = PatternFill(fill_type='solid', start_color='FFFF0000', end_color='FFFF0000')

MAXC = 199
HEADER_END = 14
BOTTOM_FIRST_SRC = 31


def parse_rows(ref):
    a, _, b = ref.partition(':')
    ra = int(''.join(ch for ch in a if ch.isdigit()))
    if not b:
        return ra, ra
    rb = int(''.join(ch for ch in b if ch.isdigit()))
    return ra, rb


def shift_ref(ref, offset):
    a, _, b = ref.partition(':')
    def sh(p):
        let = ''.join(ch for ch in p if not ch.isdigit())
        num = int(''.join(ch for ch in p if ch.isdigit()))
        return '%s%d' % (let, num + offset)
    return '%s:%s' % (sh(a), sh(b)) if b else sh(a)


def cap_cell(cell):
    return {
        'v': cell.value,
        'font': copy(cell.font),
        'border': copy(cell.border),
        'fill': copy(cell.fill),
        'alignment': copy(cell.alignment),
        'number_format': cell.number_format,
        'protection': copy(cell.protection),
    }


def apply_cell(cell, cap):
    for k in ('font', 'border', 'fill', 'alignment', 'protection'):
        setattr(cell, k, copy(cap[k]))
    cell.number_format = cap['number_format']


def cap_row(ws, r):
    return {
        'height': ws.row_dimensions[r].height if r in ws.row_dimensions else None,
        'cells': [cap_cell(ws.cell(r, c)) for c in range(1, MAXC + 1)],
    }


def apply_row(ws, r, cap, with_value):
    if cap['height'] is not None:
        ws.row_dimensions[r].height = cap['height']
    for c in range(1, MAXC + 1):
        cell = ws.cell(r, c)
        cc = cap['cells'][c - 1]
        apply_cell(cell, cc)
        if with_value:
            cell.value = cc['v']


def main():
    tpl, datafile, out = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(datafile, 'r', encoding='utf-8') as f:
        D = json.load(f)

    wb = load_workbook(tpl)
    ws = wb.active
    last = ws.max_row

    # Стили/значения берём ДО перестройки
    st_gar = cap_row(ws, 15)
    st_dept = cap_row(ws, 16)
    st_sub = cap_row(ws, 18)
    st_grand = cap_row(ws, 26)
    bottom = [cap_row(ws, r) for r in range(BOTTOM_FIRST_SRC, last + 1)]

    merges = [str(r) for r in ws.merged_cells.ranges]
    header_merges = [r for r in merges if parse_rows(r)[1] <= HEADER_END]
    bottom_merges = [r for r in merges if parse_rows(r)[0] >= BOTTOM_FIRST_SRC + 1]

    # Полностью очищаем объединения и удаляем старые строки данных+шапки
    ws.merged_cells.ranges = []
    ws.delete_rows(HEADER_END + 1, last - HEADER_END)
    # openpyxl delete_rows НЕ чистит высоты устаревших строк — сбросим их вручную
    for r in list(ws.row_dimensions.keys()):
        if r and r >= HEADER_END + 1:
            del ws.row_dimensions[r]
    for ref in header_merges:
        ws.merge_cells(ref)

    rows_def = D['data_rows']
    row = HEADER_END + 1
    for dr in rows_def:
        t = dr['type']
        st = {'garrison': st_gar, 'dept': st_dept,
              'subtotal': st_sub, 'grand': st_grand}[t]
        apply_row(ws, row, st, False)
        if dr.get('text') is not None:
            ws.cell(row, 1).value = dr['text']
        if dr.get('col2') is not None:
            ws.cell(row, 2).value = dr['col2']
        for k, v in (dr.get('values') or {}).items():
            ws.cell(row, int(k)).value = v
        # Строка подразделения, у которой нет строевой или она черновик — красная заливка
        if dr.get('red'):
            for c in range(3, MAXC + 1):
                ws.cell(row, c).fill = copy(RED_FILL)
        row += 1

    N = len(rows_def)
    # Объединения строк данных
    for i, dr in enumerate(rows_def):
        r = HEADER_END + 1 + i
        if dr['type'] == 'garrison':
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=MAXC)
        elif dr['type'] in ('subtotal', 'grand'):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)

    # Нижняя шапка со сдвигом
    off = (HEADER_END + 1 + N) - BOTTOM_FIRST_SRC
    for bi, cap in enumerate(bottom):
        r = HEADER_END + 1 + N + bi
        apply_row(ws, r, cap, True)
        # подставляем должности
        for c in range(1, MAXC + 1):
            cell = ws.cell(r, c)
            val = cell.value
            if isinstance(val, str):
                cell.value = (
                    val.replace('%ДОЛЖНОСТЬ_И_ФИО_НАЧ_ДЕЖ_СМЕНЫ%', D.get('officers', {}).get('nach', '') or '')
                    .replace('%ДОЛЖНОСТЬ_И_ФИО_СТ_ПОМОЩНИКА%', D.get('officers', {}).get('st', '') or '')
                    .replace('%ДОЛЖНОСТЬ_И_ФИО_ПОМОЩНИКА%', D.get('officers', {}).get('pom', '') or '')
                    .replace('%ДОЛЖНОСТЬ_И_ФИО_ДИСПЕТЧЕРА%', D.get('officers', {}).get('disp', '') or '')
                )
    for ref in bottom_merges:
        ws.merge_cells(shift_ref(ref, off))

    # Дату в шапку
    ru = D.get('date') or ''
    for r in range(1, HEADER_END + 1):
        for c in range(1, MAXC + 1):
            cell = ws.cell(r, c)
            v = cell.value
            if isinstance(v, str) and '%ТЕК_ДАТА%' in v:
                cell.value = v.replace('%ТЕК_ДАТА%', ru)

    wb.save(out)
    print('OK')


if __name__ == '__main__':
    main()